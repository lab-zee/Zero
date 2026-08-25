"""Behavioral tests for file storage and content extraction boundaries."""

from __future__ import annotations

from types import SimpleNamespace

from docx import Document
from openpyxl import Workbook
from pptx import Presentation

from src import file_extraction, storage


def test_storage_lifecycle_and_validation(monkeypatch, tmp_path):
    monkeypatch.setattr(storage, "UPLOAD_DIR", str(tmp_path))
    path = storage.save_file(b"payload", "report.txt")

    assert storage.file_exists(path)
    assert storage.get_file(path) == b"payload"
    assert storage.validate_file_type("report.txt", "text/plain; charset=utf-8") == (True, None)
    assert storage.validate_file_type("report.txt", "application/x-custom") == (True, None)
    allowed, error = storage.validate_file_type("payload.exe")
    assert not allowed
    assert "'.exe' is not allowed" in error
    assert storage.generate_unique_filename("report.PDF").endswith(".PDF")
    assert storage.delete_file_from_storage(path)
    assert not storage.delete_file_from_storage(path)


def test_storage_delete_handles_os_error(monkeypatch):
    monkeypatch.setattr(storage.os.path, "exists", lambda _path: True)
    monkeypatch.setattr(storage.os, "remove", lambda _path: (_ for _ in ()).throw(OSError("busy")))
    assert storage.delete_file_from_storage("busy.txt") is False


def test_extract_text_and_latin1_files(tmp_path):
    utf8 = tmp_path / "notes.txt"
    utf8.write_text("hello\nworld", encoding="utf-8")
    latin1 = tmp_path / "legacy.csv"
    latin1.write_bytes("café".encode("latin-1"))

    assert file_extraction.extract_text_from_file(str(utf8)) == ("hello\nworld", None)
    assert file_extraction.extract_text_from_file(str(latin1))[0] == "café"


def test_extract_missing_empty_and_unsupported_files(tmp_path):
    missing = tmp_path / "missing.txt"
    assert "File not found" in file_extraction.extract_text_from_file(str(missing))[1]

    empty = tmp_path / "empty.txt"
    empty.write_bytes(b"")
    assert "File is empty" in file_extraction.extract_text_from_file(str(empty))[1]

    unsupported = tmp_path / "archive.zip"
    unsupported.write_bytes(b"zip")
    assert "not yet supported" in file_extraction.extract_text_from_file(str(unsupported))[1]


def test_extract_excel_word_and_powerpoint(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Metrics"
    sheet.append(["Revenue", 42])
    excel = tmp_path / "metrics.xlsx"
    workbook.save(excel)

    document = Document()
    document.add_paragraph("Executive summary")
    word = tmp_path / "memo.docx"
    document.save(word)

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Roadmap"
    slide.placeholders[1].text = "Launch"
    powerpoint = tmp_path / "plan.pptx"
    presentation.save(powerpoint)

    excel_text, excel_error = file_extraction.extract_text_from_file(str(excel))
    word_text, word_error = file_extraction.extract_text_from_file(str(word))
    ppt_text, ppt_error = file_extraction.extract_text_from_file(str(powerpoint))

    assert excel_error is None and "Metrics" in excel_text and "Revenue\t42" in excel_text
    assert word_error is None and word_text == "Executive summary"
    assert ppt_error is None and "Slide 1" in ppt_text and "Roadmap" in ppt_text and "Launch" in ppt_text


def test_extract_corrupt_structured_files_reports_specific_errors(tmp_path):
    expected = {
        "broken.pdf": "Error extracting PDF",
        "broken.xlsx": "Error extracting Excel",
        "broken.docx": "Error extracting Word",
        "broken.pptx": "Error extracting PowerPoint",
    }
    for filename, message in expected.items():
        path = tmp_path / filename
        path.write_bytes(b"not a real document")
        assert message in file_extraction.extract_text_from_file(str(path))[1]


def test_extract_files_content_combines_success_and_failures(monkeypatch):
    files = [
        SimpleNamespace(original_filename="good.txt", file_path="/good", content_type="text/plain"),
        SimpleNamespace(original_filename="bad.pdf", file_path="/bad", content_type="application/pdf"),
        SimpleNamespace(original_filename="blank.txt", file_path="/blank", content_type="text/plain"),
    ]
    outcomes = {
        "/good": ("useful content", None),
        "/bad": ("", "encrypted"),
        "/blank": ("", None),
    }
    monkeypatch.setattr(
        file_extraction,
        "extract_text_from_file",
        lambda path, _content_type, _filename: outcomes[path],
    )

    combined, errors = file_extraction.extract_files_content(files)

    assert "good.txt" in combined and "useful content" in combined
    assert "bad.pdf" in combined and "EXTRACTION FAILED: encrypted" in combined
    assert errors == [
        "Error extracting bad.pdf: encrypted",
        "No content extracted from blank.txt",
    ]
    assert file_extraction.extract_files_content([]) == ("", [])
